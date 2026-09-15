# ============================================================
# NagarSathi - Municipal RAG Builder
# Reads PDF, Word, Excel and URLs from the data folders
# ============================================================

import os
import re
import shutil
from pathlib import Path

import pandas as pd
import pymupdf
import pytesseract

from PIL import Image
from docx import Document as DocxDocument

from dotenv import load_dotenv

from langchain_core.documents import Document

from langchain_community.document_loaders import (
    PyPDFLoader,
    WebBaseLoader
)

from langchain_text_splitters import RecursiveCharacterTextSplitter

from langchain_openai import OpenAIEmbeddings

from langchain_community.vectorstores import FAISS


# ============================================================
# 1. LOAD ENVIRONMENT
# ============================================================

load_dotenv()

os.environ["USER_AGENT"] = "NagarSathi/1.0"


# ============================================================
# 2. FOLDER PATHS
# ============================================================

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
PDF_DIR = DATA_DIR / "pdf_file"
WORD_DIR = DATA_DIR / "Word_file"
EXCEL_DIR = DATA_DIR / "Excel_file"
URL_FILE = DATA_DIR / "url.txt"
VECTORSTORE_DIR = BASE_DIR / "vectorstore"

# ============================================================
# 3. CREATE FOLDERS IF THEY DON'T EXIST
# ============================================================

PDF_DIR.mkdir(parents=True, exist_ok=True)
WORD_DIR.mkdir(parents=True, exist_ok=True)
EXCEL_DIR.mkdir(parents=True, exist_ok=True)


# ============================================================
# 4. CLEAN TEXT
# ============================================================

def clean_text(text):

    if not text:
        return ""

    text = text.replace("\x00", " ")

    text = re.sub(
        r"[ \t]+",
        " ",
        text
    )

    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text
    )

    return text.strip()


# ============================================================
# 5. CONFIGURE TESSERACT
# ============================================================

def configure_tesseract():

    tesseract_path = os.getenv("TESSERACT_CMD")

    if tesseract_path:

        pytesseract.pytesseract.tesseract_cmd = (
            tesseract_path
        )

        return True

    # Try common Windows installation locations

    possible_paths = [

        r"C:\Program Files\Tesseract-OCR\tesseract.exe",

        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"

    ]

    for path in possible_paths:

        if Path(path).exists():

            pytesseract.pytesseract.tesseract_cmd = path

            return True

    return False


# ============================================================
# 6. LOAD NORMAL / SCANNED PDF
# ============================================================

def load_pdf(file_path):

    documents = []

    print()
    print("----------------------------------------")
    print("Processing PDF:", file_path.name)
    print("----------------------------------------")

    # --------------------------------------------------------
    # First attempt: normal PDF text extraction
    # --------------------------------------------------------

    loader = PyPDFLoader(
        str(file_path)
    )

    try:

        pdf_docs = loader.load()

    except Exception as e:

        print("PDF loading error:", e)

        pdf_docs = []


    # --------------------------------------------------------
    # Check extracted text
    # --------------------------------------------------------

    total_text = sum(
        len(doc.page_content.strip())
        for doc in pdf_docs
    )


    # --------------------------------------------------------
    # NORMAL TEXT PDF
    # --------------------------------------------------------

    if total_text > 50:

        print("✓ Normal text PDF detected")

        for doc in pdf_docs:

            text = clean_text(
                doc.page_content
            )

            if not text:
                continue

            page_number = (
                doc.metadata.get("page", 0) + 1
            )

            documents.append(

                Document(

                    page_content=text,

                    metadata={
                        "source_file":
                            file_path.name,

                        "source_path":
                            str(file_path.resolve()),

                        "source_folder":
                            PDF_DIR.name,

                        "file_type":
                            "PDF",

                        "page":
                            page_number,
                    }
                )

            )

        print(
            "✓ Pages loaded:",
            len(documents)
        )

        return documents


    # --------------------------------------------------------
    # SCANNED PDF
    # --------------------------------------------------------

    print("→ Scanned PDF detected")

    if not configure_tesseract():

        print(
            "✗ Tesseract OCR is not installed."
        )

        print(
            "Install Tesseract or set "
            "TESSERACT_CMD in .env"
        )

        return documents


    pdf = pymupdf.open(
        str(file_path)
    )


    for page_number, page in enumerate(
        pdf,
        start=1
    ):

        print(
            f"  OCR page {page_number}"
        )

        pix = page.get_pixmap(
            matrix=pymupdf.Matrix(
                2,
                2
            ),
            alpha=False
        )

        image = Image.frombytes(

            "RGB",

            [
                pix.width,
                pix.height
            ],

            pix.samples

        )

        text = pytesseract.image_to_string(
            image
        )

        text = clean_text(text)

        if text:

            documents.append(

                Document(

                    page_content=text,

                    metadata={
                        "source_file":
                            file_path.name,

                        "source_path":
                            str(file_path.resolve()),

                        "source_folder":
                            PDF_DIR.name,

                        "file_type":
                            "Scanned PDF",

                        "page":
                            page_number,
                    }
                )

            )


    pdf.close()

    print(
        "✓ OCR completed:",
        len(documents),
        "pages"
    )

    return documents


# ============================================================
# 7. LOAD WORD FILE
# ============================================================

def load_word(file_path):

    print()
    print("----------------------------------------")
    print("Processing Word:", file_path.name)
    print("----------------------------------------")

    doc = DocxDocument(
        str(file_path)
    )

    parts = []


    # --------------------------------------------------------
    # Paragraphs
    # --------------------------------------------------------

    for paragraph in doc.paragraphs:

        text = clean_text(
            paragraph.text
        )

        if text:

            parts.append(text)


    # --------------------------------------------------------
    # Tables
    # --------------------------------------------------------

    for table in doc.tables:

        for row in table.rows:

            row_values = []

            for cell in row.cells:

                value = clean_text(
                    cell.text
                )

                if value:

                    row_values.append(value)


            if row_values:

                parts.append(
                    " | ".join(row_values)
                )


    full_text = "\n\n".join(
        parts
    )


    if not full_text:

        print("⚠ No text found")

        return []


    print("✓ Word loaded")


    return [

        Document(

            page_content=full_text,

            metadata={
                "source_file":
                    file_path.name,

                "source_path":
                    str(
                        file_path.resolve()
                    ),

                "source_folder":
                    WORD_DIR.name,

                "file_type":
                    "Word"

            }

        )

    ]


# ============================================================
# 8. LOAD EXCEL
# ============================================================

def load_excel(file_path):

    documents = []

    print()
    print("----------------------------------------")
    print("Processing Excel:", file_path.name)
    print("----------------------------------------")


    # --------------------------------------------------------
    # Read workbook
    # --------------------------------------------------------

    try:

        excel_file = pd.ExcelFile(
            file_path
        )

    except Exception as e:

        print(
            "✗ Excel error:",
            e
        )

        return documents


    # --------------------------------------------------------
    # Process every sheet
    # --------------------------------------------------------

    for sheet_name in excel_file.sheet_names:

        print(
            "  Sheet:",
            sheet_name
        )

        try:

            df = pd.read_excel(

                file_path,

                sheet_name=sheet_name

            )

        except Exception as e:

            print(
                "  ✗ Sheet error:",
                e
            )

            continue


        df = df.fillna("")


        rows = []


        # ----------------------------------------------------
        # Convert every row into searchable text
        # ----------------------------------------------------

        for row_number, row in df.iterrows():

            values = []


            for column, value in row.items():

                value = str(value).strip()


                if value:

                    values.append(

                        f"{column}: {value}"

                    )


            if values:

                rows.append(

                    " | ".join(values)

                )


        sheet_text = "\n".join(
            rows
        )


        if sheet_text:

            documents.append(

                Document(

                    page_content=sheet_text,

                    metadata={

                        "source_file":
                            file_path.name,

                        "source_path":
                            str(
                                file_path.resolve()
                            ),

                        "source_folder":
                            EXCEL_DIR.name,

                        "file_type":
                            "Excel",

                        "sheet":
                            sheet_name

                    }

                )

            )


    print(
        "✓ Excel sheets loaded:",
        len(documents)
    )

    return documents


# ============================================================
# 9. LOAD URLS
# ============================================================

def load_urls():

    documents = []


    if not URL_FILE.exists():

        print(
            "No url.txt file found."
        )

        return documents


    # --------------------------------------------------------
    # Read and de-duplicate configured URLs
    # --------------------------------------------------------

    urls = []
    seen_urls = set()

    try:

        configured_lines = (
            URL_FILE.read_text(
                encoding="utf-8"
            ).splitlines()
        )

    except Exception as e:

        print(
            "✗ URL file reading error:",
            e
        )

        return documents

    for line in configured_lines:

        url = str(
            line or ""
        ).strip()

        if not url:
            continue

        if not url.startswith(
            (
                "http://",
                "https://"
            )
        ):
            continue

        # Remove accidental trailing slash
        # only for duplicate comparison.
        normalized_url = url.rstrip("/").lower()

        if normalized_url in seen_urls:
            continue

        seen_urls.add(
            normalized_url
        )

        urls.append(url)


    print()
    print("----------------------------------------")
    print("Processing URLs")
    print("----------------------------------------")


    for url in urls:

        print(
            "URL:",
            url
        )


        try:

            loader = WebBaseLoader(
                url
            )

            web_docs = loader.load()


            for doc in web_docs:

                text = clean_text(
                    doc.page_content
                )

                if not text:

                    continue


                documents.append(

                    Document(

                        page_content=text,

                        metadata={

                            "source_file":
                                url,

                            "source_folder":
                                "Official URL",

                            "source_url":
                                url,

                            "file_type":
                                "Web"

                        }

                    )

                )


            print("✓ URL loaded")


        except Exception as e:

            print(
                "✗ URL failed:",
                e
            )


    return documents


# ============================================================
# 10. LOAD EVERYTHING
# ============================================================

def load_all_documents():

    documents = []


    # --------------------------------------------------------
    # PDF
    # --------------------------------------------------------

    for file in PDF_DIR.glob("*.pdf"):

        documents.extend(
            load_pdf(file)
        )


    # --------------------------------------------------------
    # WORD
    # --------------------------------------------------------

    for file in WORD_DIR.glob("*.docx"):

        documents.extend(
            load_word(file)
        )


    # --------------------------------------------------------
    # EXCEL
    # --------------------------------------------------------

    for file in EXCEL_DIR.iterdir():

        if not file.is_file():
            continue

        if file.suffix.lower() in [
            ".xlsx",
            ".xls",
            ".xlsm"
        ]:

            documents.extend(
                load_excel(file)
            )

    # --------------------------------------------------------
    # URL
    # --------------------------------------------------------

    documents.extend(
        load_urls()
    )


    return documents


# ============================================================
# 11. BUILD VECTOR DATABASE
# ============================================================

def build_vectorstore():

    print()
    print("========================================")
    print("NAGARSATHI RAG BUILD")
    print("========================================")

    try:

        documents = load_all_documents()

    except Exception as e:

        print()
        print(
            "✗ Document loading failed:",
            e
        )

        return False


    # --------------------------------------------------------
    # Check documents
    # --------------------------------------------------------

    if not documents:

        print()
        print(
            "⚠ No documents found."
        )

        print(
            "Please place files inside:"
        )

        print(
            "data/pdf_file"
        )

        print(
            "data/Word_file"
        )

        print(
            "data/Excel_file"
        )

        print(
            "and URLs inside:"
        )

        print(
            "data/url.txt"
        )

        return False


    print()
    print(
        "Total source documents:",
        len(documents)
    )


    # ========================================================
    # 12. SPLIT DOCUMENTS
    # ========================================================

    text_splitter = (
        RecursiveCharacterTextSplitter(

            chunk_size=1000,

            chunk_overlap=150,

            separators=[

                "\n\n",

                "\n",

                ". ",

                " ",

                ""

            ]

        )
    )


    chunks = text_splitter.split_documents(
        documents
    )


    print(
        "Total chunks:",
        len(chunks)
    )


    # ========================================================
    # 13. CREATE EMBEDDINGS
    # ========================================================

    print()
    print(
        "Creating embeddings..."
    )

    try:

        embeddings = OpenAIEmbeddings(
            model="text-embedding-3-small"
        )

    except Exception as e:

        print()
        print(
            "✗ Embedding initialization failed:",
            e
        )

        return False


    # ========================================================
    # 14. CREATE NEW FAISS FIRST
    #
    # IMPORTANT:
    # Do NOT delete the existing vectorstore until the new
    # vectorstore has been successfully created.
    # ========================================================

    try:

        vectorstore = FAISS.from_documents(
            chunks,
            embeddings
        )

    except Exception as e:

        print()
        print(
            "✗ FAISS creation failed:",
            e
        )

        return False


    # ========================================================
    # 15. SAVE NEW VECTORSTORE
    # ========================================================

    temporary_vectorstore_dir = (
        VECTORSTORE_DIR.parent
        / (
            VECTORSTORE_DIR.name
            + "_new"
        )
    )

    try:

        # Remove an incomplete previous temporary build.
        if temporary_vectorstore_dir.exists():

            shutil.rmtree(
                temporary_vectorstore_dir
            )

        temporary_vectorstore_dir.mkdir(
            parents=True,
            exist_ok=True
        )

        vectorstore.save_local(
            str(
                temporary_vectorstore_dir
            )
        )

    except Exception as e:

        print()
        print(
            "✗ Vectorstore save failed:",
            e
        )

        if temporary_vectorstore_dir.exists():

            try:
                shutil.rmtree(
                    temporary_vectorstore_dir
                )
            except Exception:
                pass

        return False


    # ========================================================
    # 16. REPLACE OLD VECTORSTORE
    #
    # Only replace the old database after the new one has
    # been completely written.
    # ========================================================

    try:

        if VECTORSTORE_DIR.exists():

            shutil.rmtree(
                VECTORSTORE_DIR
            )

        temporary_vectorstore_dir.rename(
            VECTORSTORE_DIR
        )

    except Exception as e:

        print()
        print(
            "✗ Vectorstore replacement failed:",
            e
        )

        return False

    print()
    print("========================================")
    print("✓ VECTOR DATABASE CREATED")
    print("========================================")

    print(
        "Source documents:",
        len(documents)
    )

    print(
        "Chunks:",
        len(chunks)
    )

    print(
        "Saved to:",
        VECTORSTORE_DIR
    )


    return True


# ============================================================
# 17. RUN
# ============================================================

if __name__ == "__main__":

    build_vectorstore()
